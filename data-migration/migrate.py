import openpyxl
import re
from datetime import datetime
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
import uuid

# Database setup
DATABASE_URL = "postgresql://root:password@localhost:5432/diabetes_db"
engine = create_engine(DATABASE_URL)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
session = SessionLocal()

# Default admin user id
admin_id = str(uuid.uuid4())

try:
    session.execute(text('INSERT INTO "User" (id, username, password, "updatedAt") VALUES (:id, \'admin\', \'admin\', now()) ON CONFLICT (username) DO NOTHING'), {"id": admin_id})
    session.commit()
    res = session.execute(text("SELECT id FROM \"User\" WHERE username = 'admin'"))
    admin_id = res.scalar()
except Exception as e:
    session.rollback()
    print("Error with admin user:", e)

# CLEAR OLD DATA
session.execute(text('TRUNCATE "Prescription", "DispensingRule", "Item" CASCADE'))
session.commit()

# Read excel with openpyxl to get values and colors
wb = openpyxl.load_workbook('/Users/drpadrino/Documents/WorkIaTdf/control diabetes 26.xlsx', data_only=True)
sheet = wb['Hoja1']

header_row = 3
items_map = {} # col_idx -> item_id
current_item_name = ""

for col_idx in range(3, 25): # C to X
    item_name = sheet.cell(row=header_row, column=col_idx).value
    if item_name and str(item_name).strip() != "" and str(item_name).strip() != 'OBSERVACION':
        current_item_name = str(item_name).strip()
        
    if current_item_name:
        item_id = str(uuid.uuid4())
        session.execute(text('INSERT INTO "Item" (id, name, "updatedAt") VALUES (:id, :name, now()) ON CONFLICT (name) DO NOTHING'), {"id": item_id, "name": current_item_name})
        res = session.execute(text('SELECT id FROM "Item" WHERE name = :name'), {"name": current_item_name})
        item_id = res.scalar()
        items_map[col_idx] = item_id
session.commit()

pattern = re.compile(r'\((\d+)\)\s*(\d{1,2}/\d{1,2})')

for row_idx in range(4, sheet.max_row + 1):
    patient_name = sheet.cell(row=row_idx, column=1).value
    dni_val = sheet.cell(row=row_idx, column=2).value
    
    if not patient_name or not dni_val:
        continue
        
    dni = str(dni_val).strip()
    name_str = str(patient_name).strip()
    
    patient_id = str(uuid.uuid4())
    session.execute(text('INSERT INTO "Patient" (id, dni, "firstName", "lastName", "patientType", "updatedAt") VALUES (:id, :dni, \'\', :lastName, \'TIPO_2\', now()) ON CONFLICT (dni) DO NOTHING'), {"id": patient_id, "dni": dni, "lastName": name_str})
    res = session.execute(text('SELECT id FROM "Patient" WHERE dni = :dni'), {"dni": dni})
    patient_id = res.scalar()
    
    for col_idx, item_id in items_map.items():
        cell = sheet.cell(row=row_idx, column=col_idx)
        val = cell.value
        if val:
            val_str = str(val).strip()
            matches = pattern.findall(val_str)
            for match in matches:
                qty = int(match[0])
                date_str = match[1] + "/2026"
                try:
                    date_obj = datetime.strptime(date_str, "%d/%m/%Y")
                except:
                    continue
                
                status = "AUTHORIZED"
                if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
                    color = str(cell.fill.fgColor.rgb).upper()
                    if "FF0000" in color or "RED" in color:
                        status = "DENIED"
                    elif "FFC0CB" in color or "PINK" in color or "FF99CC" in color:
                        status = "PARTIAL"
                        
                session.execute(text("""
                    INSERT INTO "Prescription" (id, "patientId", "itemId", "datePrescribed", "quantityPrescribed", "quantityAuthorized", status, "createdById", "updatedAt") 
                    VALUES (:id, :patientId, :itemId, :date, :qty, :qty_auth, :status, :createdById, now())
                """), {
                    "id": str(uuid.uuid4()),
                    "patientId": patient_id,
                    "itemId": item_id,
                    "date": date_obj,
                    "qty": qty,
                    "qty_auth": qty if status == "AUTHORIZED" else (qty // 2 if status == "PARTIAL" else 0),
                    "status": status,
                    "createdById": admin_id
                })
session.commit()
print("Migration completed successfully.")
